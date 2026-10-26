"""Exportación a Excel (con formato, tablas y gráficos) y a CSV (para Power BI)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from cve_explorer import __version__

# Paleta del reporte
AZUL = "1B2B48"
AZUL_CLARO = "E8EEF7"
GRIS = "6B7280"
COLOR_NIVEL = {
    "INMEDIATA": "C0392B",
    "URGENTE": "E67E22",
    "PLANIFICADA": "F1C40F",
    "RUTINA": "BDC3C7",
}
COLOR_SEVERIDAD = {
    "CRÍTICA": "C0392B",
    "ALTA": "E67E22",
    "MEDIA": "F1C40F",
    "BAJA": "27AE60",
    "NINGUNA": "95A5A6",
    "SIN CALIFICAR": "BDC3C7",
}

ANCHOS = {
    "cve_id": 17, "fuente": 15, "id_fuente": 21, "titulo": 50, "descripcion": 90,
    "severidad": 14, "severidad_orden": 9, "cvss_puntaje": 9, "cvss_version": 9,
    "cvss_vector": 46, "fecha_publicacion": 13, "fecha_actualizacion": 13,
    "fecha_limite_cisa": 13, "anio": 7, "mes": 9, "proveedor": 22, "producto": 28,
    "ecosistema": 14, "cwe": 22, "explotada_activamente": 11, "ransomware": 12,
    "en_nvd": 8, "en_kev": 8, "en_github": 9, "num_fuentes": 9, "fuentes": 32,
    "estado": 24, "url": 45, "epss": 9, "epss_percentil": 10,
    "prioridad_puntaje": 10, "prioridad_nivel": 14, "prioridad_plazo": 16, "prioridad_motivo": 60,
    "activos_afectados": 30, "activo": 28, "criticidad_activo": 11, "version_instalada": 14,
    "explotadas": 11, "prioridad_max": 12,
}
ENCABEZADOS = {
    "cve_id": "CVE", "fuente": "Fuente", "id_fuente": "ID en la fuente", "titulo": "Título",
    "descripcion": "Descripción", "severidad": "Severidad", "severidad_orden": "Orden sev.",
    "cvss_puntaje": "CVSS", "cvss_version": "Versión CVSS", "cvss_vector": "Vector CVSS",
    "fecha_publicacion": "Publicada", "fecha_actualizacion": "Actualizada",
    "fecha_limite_cisa": "Límite CISA", "anio": "Año", "mes": "Mes", "proveedor": "Proveedor",
    "producto": "Producto", "ecosistema": "Ecosistema", "cwe": "CWE",
    "explotada_activamente": "Explotada", "ransomware": "Ransomware", "en_nvd": "En NVD",
    "en_kev": "En KEV", "en_github": "En GitHub", "num_fuentes": "Nº fuentes",
    "fuentes": "Fuentes", "estado": "Estado / acción", "url": "Enlace", "epss": "EPSS",
    "epss_percentil": "Percentil EPSS", "prioridad_puntaje": "Prioridad (0-100)",
    "prioridad_nivel": "Nivel", "prioridad_plazo": "Corregir en", "prioridad_motivo": "Motivo",
    "activos_afectados": "Activos afectados", "activo": "Activo", "criticidad_activo": "Criticidad",
    "version_instalada": "Versión instalada", "explotadas": "Explotadas", "prioridad_max": "Prioridad máx.",
}
COLUMNAS_PRIORIDAD = [
    "prioridad_puntaje", "prioridad_nivel", "prioridad_plazo", "cve_id", "titulo", "severidad",
    "cvss_puntaje", "epss", "epss_percentil", "explotada_activamente", "ransomware", "fecha_limite_cisa",
    "proveedor", "producto", "activos_afectados", "prioridad_motivo", "fecha_publicacion", "url",
]
COLUMNAS_BOOL = ["explotada_activamente", "en_nvd", "en_kev", "en_github"]
PELIGROSOS = ("=", "+", "-", "@", "\t", "\r")


# ----------------------------------------------------------------- limpieza

def _texto_seguro(valor, maximo: int = 32000):
    """Quita caracteres de control y evita la 'inyección de fórmulas' (CSV/Formula injection).

    Un texto que empieza con =, +, -, @ podría ejecutarse como fórmula al abrirse
    en Excel. Por eso se le antepone un espacio.
    """
    if not isinstance(valor, str):
        return valor
    valor = ILLEGAL_CHARACTERS_RE.sub("", valor)
    if valor.startswith(PELIGROSOS):
        valor = " " + valor
    return valor[:maximo]


def limpiar(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for col in df.columns:
        if df[col].dtype == object:
            df[col] = df[col].map(_texto_seguro)
    return df


def _nombre_base(desde: date, hasta: date) -> str:
    return f"cve_{desde:%Y%m%d}_{hasta:%Y%m%d}_{datetime.now():%Y%m%d-%H%M%S}"


# ----------------------------------------------------------------- CSV

def exportar_csv(tablas: dict[str, pd.DataFrame], carpeta: Path) -> list[Path]:
    """Guarda cada tabla como CSV UTF-8 (con BOM, para que Excel lea bien las tildes).

    La carpeta 'powerbi' se sobrescribe en cada corrida: así Power BI siempre
    apunta a la misma ruta y basta con darle a 'Actualizar'.
    """
    carpeta.mkdir(parents=True, exist_ok=True)
    rutas = []
    for nombre, df in tablas.items():
        ruta = carpeta / f"{nombre}.csv"
        salida = limpiar(df)
        for col in salida.columns:
            if pd.api.types.is_datetime64_any_dtype(salida[col]):
                salida[col] = salida[col].dt.strftime("%Y-%m-%d")
        salida.to_csv(ruta, index=False, encoding="utf-8-sig")
        rutas.append(ruta)
    return rutas


# ----------------------------------------------------------------- Excel

def _hoja_de_datos(writer, df: pd.DataFrame, hoja: str, nombre_tabla: str, estilo="TableStyleMedium2"):
    df = limpiar(df)
    for col in COLUMNAS_BOOL:
        if col in df.columns:
            df[col] = df[col].map({True: "Sí", False: "No"})
    columnas = list(df.columns)
    df = df.rename(columns={c: ENCABEZADOS.get(c, c) for c in columnas})
    df.to_excel(writer, sheet_name=hoja, index=False)
    ws = writer.sheets[hoja]

    for i, col in enumerate(columnas, start=1):
        letra = get_column_letter(i)
        ws.column_dimensions[letra].width = ANCHOS.get(col, 14)
        if col.startswith("fecha"):
            for celda in ws[letra][1:]:
                celda.number_format = "yyyy-mm-dd"
        elif col in ("cvss_puntaje", "prioridad_puntaje", "prioridad_max"):
            for celda in ws[letra][1:]:
                celda.number_format = "0.0"
        elif col in ("epss", "epss_percentil"):
            for celda in ws[letra][1:]:
                celda.number_format = "0.00%"

    ws.freeze_panes = "B2"
    if len(df):  # una tabla de Excel necesita al menos una fila de datos
        rango = f"A1:{get_column_letter(len(columnas))}{len(df) + 1}"
        tabla = Table(displayName=nombre_tabla, ref=rango)
        tabla.tableStyleInfo = TableStyleInfo(name=estilo, showRowStripes=True)
        ws.add_table(tabla)

        for col_color, mapa in (("severidad", COLOR_SEVERIDAD), ("prioridad_nivel", COLOR_NIVEL)):
            if col_color not in columnas:
                continue
            letra = get_column_letter(columnas.index(col_color) + 1)
            zona = f"{letra}2:{letra}{len(df) + 1}"
            for valor, color in mapa.items():
                blanco = valor in ("CRÍTICA", "ALTA", "INMEDIATA", "URGENTE")
                ws.conditional_formatting.add(
                    zona,
                    CellIsRule(operator="equal", formula=[f'"{valor}"'],
                               fill=PatternFill("solid", fgColor=color),
                               font=Font(bold=True, color="FFFFFF" if blanco else "1F2937")),
                )
    return ws


def _tabla_simple(writer, df: pd.DataFrame, hoja: str, nombre_tabla: str, anchos: list[int], formatos=None):
    df.to_excel(writer, sheet_name=hoja, index=False)
    ws = writer.sheets[hoja]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for col, formato in (formatos or {}).items():
        letra = get_column_letter(list(df.columns).index(col) + 1)
        for celda in ws[letra][1:]:
            celda.number_format = formato
    if len(df):
        tabla = Table(displayName=nombre_tabla, ref=f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}")
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tabla)
    ws.freeze_panes = "A2"
    return ws


def _hoja_resumen(wb, kpis: dict, parametros: dict, hojas: dict):
    ws = wb.create_sheet("Resumen", 0)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for col, ancho in zip("ABCDEFGHIJKLMNOP", [2, 22, 14, 2, 22, 14, 2, 22, 14, 2, 22, 14, 2, 14, 14, 14]):
        ws.column_dimensions[col].width = ancho

    ws.merge_cells("B2:L2")
    ws["B2"] = "Reporte de vulnerabilidades CVE"
    ws["B2"].font = Font(size=20, bold=True, color=AZUL)
    ws.merge_cells("B3:L3")
    ws["B3"] = (f"Periodo: {parametros['Desde']} a {parametros['Hasta']}   ·   "
                f"Fuentes: {parametros['Fuentes consultadas']}   ·   Generado: {parametros['Generado']}")
    ws["B3"].font = Font(size=10, color=GRIS)

    tarjetas = [
        ("CVE únicas", kpis["CVE únicas"], AZUL),
        ("Críticas", kpis["Críticas"], COLOR_SEVERIDAD["CRÍTICA"]),
        ("Altas", kpis["Altas"], COLOR_SEVERIDAD["ALTA"]),
        ("Explotadas activamente", kpis["Explotadas activamente"], "7D3C98"),
        ("Usadas en ransomware", kpis["Usadas en ransomware"], "922B21"),
        ("CVSS promedio", kpis["CVSS promedio"] if kpis["CVSS promedio"] is not None else "—", "2E86C1"),
        ("Prioridad INMEDIATA", kpis.get("Prioridad inmediata", "—"), COLOR_NIVEL["INMEDIATA"]),
        ("Afectan al inventario", kpis.get("Afectan al inventario", "—"), "1B5E9C"),
    ]
    orden = [0, 1, 2, 5, 3, 4, 6, 7]  # fila 1: volumen/gravedad · fila 2: explotación/prioridad
    tarjetas = [tarjetas[i] for i in orden]
    posiciones = [("B", "C"), ("E", "F"), ("H", "I"), ("K", "L")]
    for i, (titulo, valor, color) in enumerate(tarjetas):
        fila = 5 + (i // 4) * 3
        c1, c2 = posiciones[i % 4]
        ws.merge_cells(f"{c1}{fila}:{c2}{fila}")
        ws.merge_cells(f"{c1}{fila + 1}:{c2}{fila + 1}")
        ws[f"{c1}{fila}"] = titulo
        ws[f"{c1}{fila}"].font = Font(size=10, color=GRIS, bold=True)
        ws[f"{c1}{fila + 1}"] = valor
        ws[f"{c1}{fila + 1}"].font = Font(size=22, bold=True, color=color)
        ws[f"{c1}{fila + 1}"].alignment = Alignment(horizontal="left")
        for f in (fila, fila + 1):
            ws[f"{c1}{f}"].border = Border(left=Side(style="thick", color=color))
            for c in (c1, c2):
                ws[f"{c}{f}"].fill = PatternFill("solid", fgColor=AZUL_CLARO)

    fila = 12
    ws[f"B{fila}"] = "Registros por fuente"
    ws[f"B{fila}"].font = Font(bold=True, color=AZUL)
    for j, (fuente, n) in enumerate(kpis["Registros por fuente"].items(), start=1):
        ws[f"B{fila + j}"] = fuente
        ws[f"C{fila + j}"] = n
        ws[f"C{fila + j}"].number_format = "#,##0"

    # --- Gráficos (leen las hojas de tablas)
    mes = hojas["Por_Mes"]
    filas_mes = mes.max_row
    if filas_mes > 1:
        g = LineChart()
        g.title = "CVE por mes y por fuente"
        g.height, g.width = 8, 26
        g.y_axis.title = "Cantidad"
        g.y_axis.majorGridlines = None
        datos = Reference(mes, min_col=2, max_col=mes.max_column, min_row=1, max_row=filas_mes)
        g.add_data(datos, titles_from_data=True)
        g.set_categories(Reference(mes, min_col=1, min_row=2, max_row=filas_mes))
        for serie in g.series:
            serie.smooth = False
        g.x_axis.number_format = "@"
        g.x_axis.delete = False
        g.y_axis.delete = False
        ws.add_chart(g, "B17")

    sev = hojas["Por_Severidad"]
    if sev.max_row > 1:
        d = DoughnutChart()
        d.title = "Distribución por severidad"
        d.height, d.width = 8, 12
        d.add_data(Reference(sev, min_col=2, min_row=1, max_row=sev.max_row), titles_from_data=True)
        d.set_categories(Reference(sev, min_col=1, min_row=2, max_row=sev.max_row))
        serie = d.series[0]
        for idx in range(sev.max_row - 1):
            nombre = sev.cell(row=idx + 2, column=1).value
            punto = DataPoint(idx=idx)
            punto.graphicalProperties.solidFill = COLOR_SEVERIDAD.get(nombre, "999999")
            serie.dPt.append(punto)
        d.dataLabels = DataLabelList()
        d.dataLabels.showPercent = True
        d.dataLabels.showVal = False
        d.dataLabels.showCatName = False
        d.dataLabels.showSerName = False
        d.dataLabels.showLeaderLines = False
        ws.add_chart(d, "B34")

    prov = hojas["Top_Proveedores"]
    if prov.max_row > 1:
        b = BarChart()
        b.type = "bar"
        b.title = "Proveedores con más CVE"
        b.height, b.width = 8, 13.5
        b.add_data(Reference(prov, min_col=2, min_row=1, max_row=min(prov.max_row, 11)), titles_from_data=True)
        b.set_categories(Reference(prov, min_col=1, min_row=2, max_row=min(prov.max_row, 11)))
        b.y_axis.majorGridlines = None
        b.x_axis.scaling.orientation = "maxMin"  # el mayor arriba
        b.legend = None
        b.x_axis.delete = False
        b.y_axis.delete = False
        b.series[0].graphicalProperties.solidFill = AZUL
        ws.add_chart(b, "H34")

    cwe = hojas["Top_CWE"]
    if cwe.max_row > 1:
        b = BarChart()
        b.type = "bar"
        b.title = "Debilidades (CWE) más comunes"
        b.height, b.width = 8, 26
        b.add_data(Reference(cwe, min_col=3, min_row=1, max_row=min(cwe.max_row, 11)), titles_from_data=True)
        b.set_categories(Reference(cwe, min_col=1, min_row=2, max_row=min(cwe.max_row, 11)))
        b.y_axis.majorGridlines = None
        b.x_axis.scaling.orientation = "maxMin"
        b.legend = None
        b.x_axis.delete = False
        b.y_axis.delete = False
        b.series[0].graphicalProperties.solidFill = "2E86C1"
        ws.add_chart(b, "B51")
    return ws


def exportar_excel(
    ruta: Path,
    registros: pd.DataFrame,
    consolidado: pd.DataFrame,
    resumenes: dict[str, pd.DataFrame],
    kpis: dict,
    parametros: dict,
) -> Path:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        top = resumenes.get("prioridad")
        if top is not None:
            cols = [c for c in COLUMNAS_PRIORIDAD if c in top.columns]
            _hoja_de_datos(writer, top[cols], "Prioridad", "tblPrioridad", estilo="TableStyleMedium3")
        if resumenes.get("afectaciones") is not None:
            _hoja_de_datos(writer, resumenes["por_activo"], "Activos", "tblActivos", estilo="TableStyleMedium6")
            _hoja_de_datos(writer, resumenes["afectaciones"], "Inventario", "tblInventario", estilo="TableStyleMedium6")
        _hoja_de_datos(writer, consolidado, "Consolidado", "tblConsolidado")
        _hoja_de_datos(writer, registros, "Registros", "tblRegistros", estilo="TableStyleMedium9")
        hojas = {
            "Por_Mes": _tabla_simple(writer, resumenes["por_mes"], "Por_Mes", "tblPorMes",
                                     [10] + [16] * (len(resumenes["por_mes"].columns) - 1)),
            "Por_Severidad": _tabla_simple(writer, resumenes["por_severidad"], "Por_Severidad", "tblSeveridad",
                                           [16, 12, 12], {"Porcentaje": "0.0%"}),
            "Top_Proveedores": _tabla_simple(writer, resumenes["top_proveedores"], "Top_Proveedores",
                                             "tblProveedores", [30, 10, 10, 12]),
            "Top_CWE": _tabla_simple(writer, resumenes["top_cwe"], "Top_CWE", "tblCWE", [12, 48, 10]),
        }
        param = pd.DataFrame({"Parámetro": list(parametros.keys()), "Valor": [str(v) for v in parametros.values()]})
        _tabla_simple(writer, param, "Parametros", "tblParametros", [28, 100])

        wb = writer.book
        _hoja_resumen(wb, kpis, parametros, hojas)
        wb.active = 0
        wb.properties.title = "Reporte de vulnerabilidades CVE"
        wb.properties.creator = f"CVE Explorer {__version__}"
    return ruta


def exportar(
    carpeta_salida: str | Path,
    formato: str,
    desde: date,
    hasta: date,
    registros: pd.DataFrame,
    consolidado: pd.DataFrame,
    resumenes: dict[str, pd.DataFrame],
    kpis: dict,
    parametros: dict,
) -> list[Path]:
    carpeta = Path(carpeta_salida)
    rutas: list[Path] = []
    if formato in ("excel", "ambos"):
        rutas.append(exportar_excel(carpeta / f"{_nombre_base(desde, hasta)}.xlsx",
                                    registros, consolidado, resumenes, kpis, parametros))
    if formato in ("csv", "ambos"):
        tablas = {
            "consolidado": consolidado,
            **({"inventario_afectaciones": resumenes["afectaciones"]}
               if resumenes.get("afectaciones") is not None else {}),
            "registros": registros,
            "cwe_por_cve": resumenes["cwe_por_cve"],
            "por_mes": resumenes["por_mes"],
        }
        rutas.extend(exportar_csv(tablas, carpeta / "powerbi"))
    return rutas
