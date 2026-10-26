"""Exportación a Excel (con formato, tablas y gráficos) y a CSV (para Power BI)."""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.image import Image as ImagenExcel
from openpyxl.drawing.text import CharacterProperties, Font as FuenteGrafico, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from vigia import __version__
from vigia.procesamiento import serie_temporal

# Paleta del reporte
# Paleta Vigía (la misma del panel web). El naranja significa "se está explotando".
AZUL = "1B2B48"          # azul marino de la marca
AZUL_MARCA = "2A78D6"
AZUL_CLARO = "F4F6F9"
GRIS = "6B7486"
NARANJA = "EB6834"
NARANJA_TEXTO = "B8461C"
TEXTO_CLARO = "C9D6EA"   # texto secundario sobre azul marino
COLOR_NIVEL = {
    "INMEDIATA": "EB6834",
    "URGENTE": "2A78D6",
    "PLANIFICADA": "8DA9CF",
    "RUTINA": "C3CAD5",
}
# Severidad en escala de azules: más oscuro = más grave
COLOR_SEVERIDAD = {
    "CRÍTICA": "1B2B48",
    "ALTA": "2A5FA8",
    "MEDIA": "4F8FDB",
    "BAJA": "A9C8EE",
    "NINGUNA": "C9CFD8",
    "SIN CALIFICAR": "C9CFD8",
}
TEXTO_BLANCO = {"CRÍTICA", "ALTA", "MEDIA", "INMEDIATA", "URGENTE"}

ANCHOS = {
    "cve_id": 17, "fuente": 15, "id_fuente": 21, "titulo": 50, "descripcion": 90,
    "severidad": 14, "severidad_orden": 9, "cvss_puntaje": 9, "cvss_version": 9,
    "cvss_vector": 46, "fecha_publicacion": 13, "fecha_actualizacion": 13,
    "fecha_limite_cisa": 13, "anio": 7, "mes": 9, "proveedor": 22, "producto": 28,
    "ecosistema": 14, "cwe": 22, "explotada_activamente": 11, "ransomware": 12,
    "en_nvd": 8, "en_kev": 8, "en_github": 9, "num_fuentes": 9, "fuentes": 32,
    "estado": 24, "url": 45, "epss": 9, "epss_percentil": 10,
    "prioridad_puntaje": 9, "prioridad_nivel": 13, "prioridad_plazo": 16, "prioridad_motivo": 60,
    "activos_afectados": 30, "activo": 28, "criticidad_activo": 11, "version_instalada": 14,
    "explotadas": 11, "prioridad_max": 12,
}
ENCABEZADOS = {
    "cve_id": "CVE", "fuente": "Fuente", "id_fuente": "ID en la fuente", "titulo": "Título",
    "descripcion": "Descripción", "severidad": "Severidad", "severidad_orden": "Orden sev.",
    "cvss_puntaje": "CVSS", "cvss_version": "Versión CVSS", "cvss_vector": "Vector CVSS",
    "fecha_publicacion": "Publicada", "fecha_actualizacion": "Actualizada",
    "fecha_limite_cisa": "Límite CISA", "anio": "Año", "mes": "Mes", "proveedor": "Proveedor",
    "producto": "Producto", "ecosistema": "Ecosistema", "cwe": "CWE",
    "explotada_activamente": "Explotada", "ransomware": "Ransomware", "en_nvd": "En NVD",
    "en_kev": "En KEV", "en_github": "En GitHub", "num_fuentes": "Nº fuentes",
    "fuentes": "Fuentes", "estado": "Estado / acción", "url": "Enlace", "epss": "EPSS",
    "epss_percentil": "Percentil EPSS", "prioridad_puntaje": "Índice",
    "prioridad_nivel": "Nivel", "prioridad_plazo": "Corregir en", "prioridad_motivo": "Motivo",
    "activos_afectados": "Activos afectados", "activo": "Activo", "criticidad_activo": "Criticidad",
    "version_instalada": "Versión instalada", "explotadas": "Explotadas", "prioridad_max": "Prioridad máx.",
}
COLUMNAS_PRIORIDAD = [
    "prioridad_puntaje", "prioridad_nivel", "prioridad_plazo", "cve_id", "titulo", "severidad",
    "cvss_puntaje", "epss", "epss_percentil", "explotada_activamente", "ransomware", "fecha_limite_cisa",
    "proveedor", "producto", "activos_afectados", "prioridad_motivo", "fecha_publicacion", "url",
]
COLUMNAS_BOOL = ["explotada_activamente", "en_nvd", "en_kev", "en_github"]
PELIGROSOS = ("=", "+", "-", "@", "\t", "\r")
# Caracteres que XML 1.0 no permite y que openpyxl no filtra (U+FFFE, U+FFFF y
# sustitutos sueltos). Si se cuelan, Excel dice que el archivo está dañado.
NO_VALIDOS_XML = re.compile("[\ufffe\uffff\ud800-\udfff]")


# ----------------------------------------------------------------- limpieza

def _texto_seguro(valor, maximo: int = 32000):
    """Quita caracteres de control y evita la 'inyección de fórmulas' (CSV/Formula injection).

    Un texto que empieza con =, +, -, @ podría ejecutarse como fórmula al abrirse
    en Excel. Por eso se le antepone un espacio.
    """
    if not isinstance(valor, str):
        return valor
    valor = ILLEGAL_CHARACTERS_RE.sub("", valor)
    valor = NO_VALIDOS_XML.sub("", valor)
    if valor.startswith(PELIGROSOS):
        valor = " " + valor
    return valor[:maximo]


def limpiar(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for col in df.columns:
        if df[col].dtype == object:
            df[col] = df[col].map(_texto_seguro)
    return df


def _nombre_base(desde: date, hasta: date) -> str:
    return f"cve_{desde:%Y%m%d}_{hasta:%Y%m%d}_{datetime.now():%Y%m%d-%H%M%S}"


# ----------------------------------------------------------------- CSV

def exportar_csv(tablas: dict[str, pd.DataFrame], carpeta: Path) -> list[Path]:
    """Guarda cada tabla como CSV UTF-8 (con BOM, para que Excel lea bien las tildes).

    La carpeta 'powerbi' se sobrescribe en cada corrida: así Power BI siempre
    apunta a la misma ruta y basta con darle a 'Actualizar'.
    """
    carpeta.mkdir(parents=True, exist_ok=True)
    rutas = []
    for nombre, df in tablas.items():
        ruta = carpeta / f"{nombre}.csv"
        salida = limpiar(df)
        for col in salida.columns:
            if pd.api.types.is_datetime64_any_dtype(salida[col]):
                salida[col] = salida[col].dt.strftime("%Y-%m-%d")
        salida.to_csv(ruta, index=False, encoding="utf-8-sig")
        rutas.append(ruta)
    return rutas


# ----------------------------------------------------------------- Excel

# Texto de color (sin rellenos) para severidad, plazo y explotación: se lee
# de un vistazo sin llenar la hoja de bloques de color.
TEXTO_SEVERIDAD = {"Crítica": ("1B2B48", True), "Alta": ("2A5FA8", True), "Media": ("4F8FDB", False),
                   "Baja": ("6B7486", False), "Ninguna": ("9AA3B2", False), "Sin calificar": ("9AA3B2", False)}
TEXTO_NIVEL = {"Inmediata": (NARANJA_TEXTO, True), "Urgente": ("2A78D6", True),
               "Planificada": ("6B7486", False), "Rutina": ("9AA3B2", False)}
TEXTO_SI_NO = {"Sí": (NARANJA_TEXTO, True), "No": ("9AA3B2", False)}
COLUMNAS_SUAVES = ("prioridad_plazo", "fuentes", "cvss_version", "estado", "id_fuente")
# Columnas de texto que van justo después de otra larga: un poco de aire a la izquierda
COLUMNAS_SEPARADAS = ("severidad", "prioridad_nivel", "proveedor", "producto", "activos_afectados",
                      "prioridad_motivo", "cwe", "ecosistema")


def _encabezado_marca(ws, n_columnas: int) -> None:
    """Fila de títulos en azul marino con letra blanca, como la franja del panel."""
    for i in range(1, n_columnas + 1):
        celda = ws.cell(row=1, column=i)
        celda.fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
        celda.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30


def _texto_condicional(ws, columna: int, filas: int, mapa: dict) -> None:
    letra = get_column_letter(columna)
    for valor, (color, negrita) in mapa.items():
        ws.conditional_formatting.add(
            f"{letra}2:{letra}{filas + 1}",
            CellIsRule(operator="equal", formula=[f'"{valor}"'], font=Font(bold=negrita, color=color)),
        )


def _hoja_de_datos(writer, df: pd.DataFrame, hoja: str, nombre_tabla: str, estilo="TableStyleLight1"):
    df = limpiar(df)
    for col in COLUMNAS_BOOL:
        if col in df.columns:
            df[col] = df[col].map({True: "Sí", False: "No"})
    # Menos ruido: mayúsculas solo para siglas y "Desconocido" no se repite en cada fila
    for col in ("severidad", "prioridad_nivel", "criticidad_activo"):
        if col in df.columns:
            df[col] = df[col].map(lambda v: v.capitalize() if isinstance(v, str) else v)
    if "ransomware" in df.columns:
        df["ransomware"] = df["ransomware"].map(lambda v: "Sí" if v == "Sí" else None)
    columnas = list(df.columns)
    df = df.rename(columns={c: ENCABEZADOS.get(c, c) for c in columnas})
    df.to_excel(writer, sheet_name=hoja, index=False)
    ws = writer.sheets[hoja]
    ws.sheet_view.showGridLines = False
    ws.sheet_format.defaultRowHeight = 18
    gris = Font(name=FUENTE, size=10, color="6B7486")

    for i, col in enumerate(columnas, start=1):
        letra = get_column_letter(i)
        ws.column_dimensions[letra].width = ANCHOS.get(col, 14)
        formato = None
        if col.startswith("fecha"):
            formato = "yyyy-mm-dd"
        elif col in ("cvss_puntaje", "prioridad_puntaje", "prioridad_max"):
            formato = "0.0"
        elif col == "epss":
            formato = "0.0%"
        elif col == "epss_percentil":
            formato = "0%"
        alineacion = None
        if col.startswith("fecha") or col in COLUMNAS_BOOL or col in ("ransomware", "cvss_version"):
            alineacion = Alignment(horizontal="center")
        elif col in COLUMNAS_SEPARADAS:
            alineacion = Alignment(horizontal="left", indent=1)
        if formato or alineacion or col in COLUMNAS_SUAVES:
            for celda in ws[letra][1:]:
                if formato:
                    celda.number_format = formato
                if alineacion:
                    celda.alignment = alineacion
                if col in COLUMNAS_SUAVES:
                    celda.font = gris

    _encabezado_marca(ws, len(columnas))
    ws.freeze_panes = "B2"
    if len(df):  # una tabla de Excel necesita al menos una fila de datos
        rango = f"A1:{get_column_letter(len(columnas))}{len(df) + 1}"
        tabla = Table(displayName=nombre_tabla, ref=rango)
        tabla.tableStyleInfo = TableStyleInfo(name=estilo, showRowStripes=True)
        ws.add_table(tabla)

        for col, mapa in (("severidad", TEXTO_SEVERIDAD), ("prioridad_nivel", TEXTO_NIVEL),
                          ("explotada_activamente", TEXTO_SI_NO), ("en_kev", TEXTO_SI_NO),
                          ("ransomware", TEXTO_SI_NO)):
            if col in columnas:
                _texto_condicional(ws, columnas.index(col) + 1, len(df), mapa)
    return ws


def _tabla_simple(writer, df: pd.DataFrame, hoja: str, nombre_tabla: str, anchos: list[int], formatos=None):
    df.to_excel(writer, sheet_name=hoja, index=False)
    ws = writer.sheets[hoja]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for col, formato in (formatos or {}).items():
        letra = get_column_letter(list(df.columns).index(col) + 1)
        for celda in ws[letra][1:]:
            celda.number_format = formato
    if len(df):
        tabla = Table(displayName=nombre_tabla, ref=f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}")
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tabla)
    _encabezado_marca(ws, len(df.columns))
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    return ws


# ----------------------------------------------------------------- hoja Resumen

# Identidad Vigía: Inter para todo el contenido de Excel. Si la PC no la tiene
# instalada, Excel usa su fuente por defecto sin dañar nada.
FUENTE = "Inter"
LOGO_EXCEL_BLANCO = Path(__file__).parent / "plantillas" / "logo-excel-blanco.png"

COLOR_FUENTE = {"NVD": "2A78D6", "CISA KEV": "EB6834", "GitHub Advisory": "1B2B48"}
GRIS_LINEA = "E5E7EB"
ALTO_GRAFICO = 7.6      # cm (≈ 14 filas de 15 pt)
ANCHO_COMPLETO = 29.0   # cm (columnas B a L)
ANCHO_MEDIO = 14.2      # cm


def _titulo(texto: str) -> Title:
    """Título que NO se monta encima del gráfico.

    openpyxl no escribe <c:overlay val="0"/> y, sin esa etiqueta, Excel pone el
    título y la leyenda ENCIMA del área del gráfico (en LibreOffice no pasa).
    """
    letra = CharacterProperties(sz=1200, b=True, solidFill="1B2B48", latin=FuenteGrafico(typeface=FUENTE))
    parrafo = Paragraph(pPr=ParagraphProperties(defRPr=letra), r=[RegularTextRun(rPr=letra, t=texto)])
    return Title(tx=Text(rich=RichText(p=[parrafo])), overlay=False)


def _texto_ejes(eje, tamano: int = 900) -> None:
    letra = CharacterProperties(sz=tamano, solidFill="4B5563", latin=FuenteGrafico(typeface=FUENTE))
    eje.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=letra), endParaRPr=letra)])


def _lineas_suaves():
    return ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=GRIS_LINEA)))


def _etiquetas_valor() -> DataLabelList:
    """Solo el número al final de cada barra (sin nombre de serie ni de categoría)."""
    etiquetas = DataLabelList()
    etiquetas.showVal = True
    for campo in ("showSerName", "showCatName", "showLegendKey", "showPercent", "showLeaderLines"):
        setattr(etiquetas, campo, False)
    etiquetas.position = "outEnd"
    etiquetas.numFmt = "#,##0"
    return etiquetas


def _barras_horizontales(hoja, titulo: str, col_categoria: int, col_valor: int, filas: int,
                         colores: list[str] | None = None, color: str = AZUL) -> BarChart:
    """Barras horizontales limpias: categorías en orden (la mayor arriba), números al final
    de cada barra y sin eje de valores (los números ya están en las etiquetas)."""
    g = BarChart()
    g.type = "bar"
    g.title = _titulo(titulo)
    g.style = 10
    g.height, g.width = ALTO_GRAFICO, ANCHO_MEDIO
    g.add_data(Reference(hoja, min_col=col_valor, min_row=1, max_row=filas + 1), titles_from_data=True)
    g.set_categories(Reference(hoja, min_col=col_categoria, min_row=2, max_row=filas + 1))
    g.legend = None
    g.gapWidth = 55
    g.x_axis.scaling.orientation = "maxMin"     # el primero de la tabla arriba
    g.x_axis.delete = False
    g.x_axis.majorTickMark = "none"
    g.y_axis.delete = True                      # sin eje de números: van en las etiquetas
    g.y_axis.majorGridlines = None
    serie = g.series[0]
    serie.graphicalProperties.solidFill = color
    serie.graphicalProperties.line.noFill = True
    if colores:
        for idx, c in enumerate(colores):
            punto = DataPoint(idx=idx)
            punto.graphicalProperties.solidFill = c
            punto.graphicalProperties.line.noFill = True
            serie.dPt.append(punto)
    g.dataLabels = _etiquetas_valor()
    _texto_ejes(g.x_axis)
    return g


def _hoja_resumen(wb, kpis: dict, parametros: dict, hojas: dict):
    ws = wb.create_sheet("Resumen", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for col, ancho in zip("ABCDEFGHIJKLM", [2, 24, 13, 2, 24, 13, 2, 24, 13, 2, 24, 13, 2]):
        ws.column_dimensions[col].width = ancho

    # --- Franja de guardia (azul marino), igual que el panel web
    franja = PatternFill("solid", fgColor=AZUL)
    for fila in range(1, 10):
        for col in "ABCDEFGHIJKLM":
            ws[f"{col}{fila}"].fill = franja
    for fila, alto in {1: 10, 2: 36, 3: 10, 4: 20, 5: 34, 6: 34, 7: 22, 8: 20, 9: 12}.items():
        ws.row_dimensions[fila].height = alto

    logo_ok = False
    if LOGO_EXCEL_BLANCO.exists():
        try:
            logo = ImagenExcel(str(LOGO_EXCEL_BLANCO))  # necesita Pillow
            logo.height = 40
            logo.width = round(40 * 2.69)
            ws.add_image(logo, "B2")
            logo_ok = True
        except ImportError:
            pass
    if not logo_ok:
        ws["B2"] = "Vigía"
        ws["B2"].font = Font(name=FUENTE, size=18, bold=True, color="FFFFFF")
    ws.merge_cells("H2:L2")
    ws["H2"] = f"Revisado {parametros['Generado']}"
    ws["H2"].font = Font(name=FUENTE, size=10, color=TEXTO_CLARO)
    ws["H2"].alignment = Alignment(horizontal="right", vertical="center")

    ws.merge_cells("B4:G4")
    ws["B4"] = f"Parte de vigilancia del {parametros['Desde']} al {parametros['Hasta']}"
    ws["B4"].font = Font(name=FUENTE, size=11, color="9FB3D1")

    inmediatas = kpis.get("Prioridad inmediata")
    urgentes = kpis.get("Prioridad urgente")
    if isinstance(inmediatas, int) and inmediatas > 0:
        verbo = "vulnerabilidad pide" if inmediatas == 1 else "vulnerabilidades piden"
        ws.merge_cells("B5:B6")
        ws["B5"] = inmediatas
        ws["B5"].number_format = "#,##0"
        ws["B5"].font = Font(name=FUENTE, size=40, bold=True, color=NARANJA)
        ws["B5"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
        ws.merge_cells("C5:G6")
        ws["C5"] = f"{verbo} corrección inmediata."
    else:
        ws.merge_cells("B5:G6")
        ws["B5"] = "Nada urgente para corregir hoy."
    celda_titulo = ws["C5"] if isinstance(inmediatas, int) and inmediatas > 0 else ws["B5"]
    celda_titulo.font = Font(name=FUENTE, size=20, bold=True, color="FFFFFF")
    celda_titulo.alignment = Alignment(vertical="center", wrap_text=True)
    if isinstance(urgentes, int) and urgentes:
        ws.merge_cells("B7:G7")
        ws["B7"] = f"Después vienen {urgentes:,} para corregir esta semana."
        ws["B7"].font = Font(name=FUENTE, size=11, color=TEXTO_CLARO)
    ws.merge_cells("B8:G8")
    ws["B8"] = "Registros: " + ", ".join(f"{f} {n:,}" for f, n in kpis["Registros por fuente"].items())
    ws["B8"].font = Font(name=FUENTE, size=10, color="9FB3D1")

    resumen = [
        (kpis["CVE únicas"], "CVE publicadas en el periodo", "FFFFFF"),
        (kpis["Explotadas activamente"], "explotadas activamente según CISA", NARANJA),
        (kpis["Críticas"], f"críticas y {kpis['Altas']:,} altas", "FFFFFF"),
        (kpis.get("Afectan al inventario", "—"), "afectan el inventario de la organización", "FFFFFF"),
    ]
    for i, (valor, texto, color) in enumerate(resumen):
        fila = 4 + i
        ws[f"I{fila}"] = valor
        ws[f"I{fila}"].number_format = "#,##0"
        ws[f"I{fila}"].font = Font(name=FUENTE, size=15, bold=True, color=color)
        ws[f"I{fila}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
        ws.merge_cells(f"J{fila}:L{fila}")
        ws[f"J{fila}"] = texto
        ws[f"J{fila}"].font = Font(name=FUENTE, size=10, color=TEXTO_CLARO)
        ws[f"J{fila}"].alignment = Alignment(vertical="center")
        ws[f"H{fila}"].border = Border(right=Side(style="thin", color="3A4E73"))

    # --- 1) Tendencia (ancho completo)
    t = hojas["Tendencia"]
    puntos = t.max_row - 1
    if puntos >= 1 and t.max_column >= 2:
        g = LineChart()
        g.title = _titulo(f"CVE publicadas por {t['A1'].value.lower()} y por fuente")
        g.style = 10
        g.height, g.width = ALTO_GRAFICO, ANCHO_COMPLETO
        g.add_data(Reference(t, min_col=2, max_col=t.max_column, min_row=1, max_row=t.max_row), titles_from_data=True)
        g.set_categories(Reference(t, min_col=1, min_row=2, max_row=t.max_row))
        for serie in g.series:
            nombre = serie.tx.strRef.f if serie.tx and serie.tx.strRef else ""
            serie.smooth = False
            serie.marker.symbol = "circle" if puntos <= 3 else "none"
            serie.graphicalProperties.line.width = 22000
        for serie, col in zip(g.series, range(2, t.max_column + 1)):
            color = COLOR_FUENTE.get(t.cell(row=1, column=col).value, GRIS)
            serie.graphicalProperties.line.solidFill = color
            if serie.marker.symbol == "circle":
                serie.marker.graphicalProperties = GraphicalProperties(solidFill=color)
                serie.marker.graphicalProperties.line.solidFill = color
        g.legend.position = "b"
        g.legend.overlay = False
        _texto_ejes(g.x_axis)
        _texto_ejes(g.y_axis)
        g.x_axis.delete = False
        g.y_axis.delete = False
        g.x_axis.number_format = "@"
        g.x_axis.majorTickMark = "none"
        g.x_axis.tickLblSkip = max(1, puntos // 12)
        g.x_axis.tickMarkSkip = max(1, puntos // 12)
        g.y_axis.majorGridlines = _lineas_suaves()
        g.y_axis.number_format = "#,##0"
        ws.add_chart(g, "B11")

    # --- 2) Severidad y nivel de prioridad
    sev = hojas["Por_Severidad"]
    if sev.max_row > 1:
        colores = [COLOR_SEVERIDAD.get(str(sev.cell(row=i, column=1).value).upper(), "999999") for i in range(2, sev.max_row + 1)]
        ws.add_chart(_barras_horizontales(sev, "CVE por severidad (CVSS)", 1, 2, sev.max_row - 1, colores), "B28")
    niv = hojas.get("Por_Prioridad")
    if niv is not None and niv.max_row > 1:
        colores = [COLOR_NIVEL.get(str(niv.cell(row=i, column=1).value).upper(), "999999") for i in range(2, niv.max_row + 1)]
        ws.add_chart(_barras_horizontales(niv, "CVE por plazo de corrección", 1, 3, niv.max_row - 1, colores), "H28")

    # --- 3) Proveedores y debilidades
    prov = hojas["Top_Proveedores"]
    if prov.max_row > 1:
        ws.add_chart(_barras_horizontales(prov, "Proveedores con más CVE", 1, 2, min(prov.max_row - 1, 10)), "B45")
    cwe = hojas["Top_CWE"]
    if cwe.max_row > 1:
        ws.add_chart(_barras_horizontales(cwe, "Debilidades (CWE) más comunes", 4, 3, min(cwe.max_row - 1, 10),
                                          color="2A5FA8"), "H45")
    return ws


def _capitalizar(df: pd.DataFrame, columna: str) -> pd.DataFrame:
    """'CRÍTICA' -> 'Crítica' para que el gráfico no grite en mayúsculas."""
    df = df.copy()
    df[columna] = df[columna].astype(str).str.capitalize()
    return df


def _fuente_por_defecto(wb) -> None:
    """Inter en todo el libro (la fuente 0 es la que usan las celdas sin formato propio)."""
    try:
        normal = wb._fonts[0]
        wb._fonts[0] = Font(name=FUENTE, size=10, family=normal.family, scheme=None)
    except (AttributeError, IndexError, TypeError):
        pass


def _cwe_con_etiqueta(top_cwe: pd.DataFrame) -> pd.DataFrame:
    """Agrega 'Etiqueta' = 'CWE-79 · Cross-site scripting (XSS)' para el gráfico."""
    df = top_cwe.copy()
    if df.empty:
        df["Etiqueta"] = pd.Series(dtype=str)
        return df
    df["Etiqueta"] = [
        f"{c} · {d}" if d else c for c, d in zip(df["CWE"], df["Debilidad"].fillna(""))
    ]
    df["Etiqueta"] = df["Etiqueta"].map(lambda t: t if len(t) <= 42 else t[:41] + "…")
    return df


def _tabla_tendencia(registros: pd.DataFrame, desde: date, hasta: date) -> pd.DataFrame:
    """Diaria para rangos cortos (≤ 92 días) y mensual para rangos largos."""
    serie = serie_temporal(registros, desde, hasta)
    columna = "Día" if serie["granularidad"] == "día" else "Mes"
    df = pd.DataFrame({columna: serie["etiquetas"], **serie["series"]})
    if columna == "Día":
        df[columna] = pd.to_datetime(df[columna]).map(lambda d: f"{d.day}/{d.month}")  # 2026-10-05 -> 5/10
    return df


def exportar_excel(
    ruta: Path,
    registros: pd.DataFrame,
    consolidado: pd.DataFrame,
    resumenes: dict[str, pd.DataFrame],
    kpis: dict,
    parametros: dict,
    desde: date | None = None,
    hasta: date | None = None,
) -> Path:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    desde = desde or date.fromisoformat(str(parametros["Desde"]))
    hasta = hasta or date.fromisoformat(str(parametros["Hasta"]))
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        top = resumenes.get("prioridad")
        if top is not None:
            cols = [c for c in COLUMNAS_PRIORIDAD if c in top.columns]
            _hoja_de_datos(writer, top[cols], "Prioridad", "tblPrioridad", estilo="TableStyleLight1")
        if resumenes.get("afectaciones") is not None:
            _hoja_de_datos(writer, resumenes["por_activo"], "Activos", "tblActivos")
            _hoja_de_datos(writer, resumenes["afectaciones"], "Inventario", "tblInventario", estilo="TableStyleLight1")
        _hoja_de_datos(writer, consolidado, "Consolidado", "tblConsolidado")
        _hoja_de_datos(writer, registros, "Registros", "tblRegistros", estilo="TableStyleLight1")
        hojas = {
            "Por_Mes": _tabla_simple(writer, resumenes["por_mes"], "Por_Mes", "tblPorMes",
                                     [10] + [16] * (len(resumenes["por_mes"].columns) - 1)),
            "Por_Severidad": _tabla_simple(writer, _capitalizar(resumenes["por_severidad"], "Severidad"), "Por_Severidad", "tblSeveridad",
                                           [16, 12, 12], {"Porcentaje": "0.0%"}),
            "Top_Proveedores": _tabla_simple(writer, resumenes["top_proveedores"], "Top_Proveedores",
                                             "tblProveedores", [30, 10, 10, 12]),
            "Top_CWE": _tabla_simple(writer, _cwe_con_etiqueta(resumenes["top_cwe"]), "Top_CWE", "tblCWE",
                                     [12, 44, 10, 48]),
            "Tendencia": _tabla_simple(writer, _tabla_tendencia(registros, desde, hasta), "Tendencia",
                                       "tblTendencia", [12, 14, 14, 16]),
        }
        if resumenes.get("niveles") is not None:
            hojas["Por_Prioridad"] = _tabla_simple(writer, _capitalizar(resumenes["niveles"], "Nivel"), "Por_Prioridad", "tblPrioridadNiveles",
                                                   [16, 26, 10])
        param = pd.DataFrame({"Parámetro": list(parametros.keys()), "Valor": [str(v) for v in parametros.values()]})
        _tabla_simple(writer, param, "Parametros", "tblParametros", [28, 100])

        wb = writer.book
        _fuente_por_defecto(wb)
        _hoja_resumen(wb, kpis, parametros, hojas)
        wb.active = 0
        wb.properties.title = "Vigía · Reporte de vulnerabilidades CVE"
        wb.properties.creator = f"Vigía {__version__}"
    return ruta


def exportar(
    carpeta_salida: str | Path,
    formato: str,
    desde: date,
    hasta: date,
    registros: pd.DataFrame,
    consolidado: pd.DataFrame,
    resumenes: dict[str, pd.DataFrame],
    kpis: dict,
    parametros: dict,
) -> list[Path]:
    carpeta = Path(carpeta_salida)
    rutas: list[Path] = []
    if formato in ("excel", "ambos"):
        rutas.append(exportar_excel(carpeta / f"{_nombre_base(desde, hasta)}.xlsx",
                                    registros, consolidado, resumenes, kpis, parametros, desde, hasta))
    if formato in ("csv", "ambos"):
        tablas = {
            "consolidado": consolidado,
            **({"inventario_afectaciones": resumenes["afectaciones"]}
               if resumenes.get("afectaciones") is not None else {}),
            "registros": registros,
            "cwe_por_cve": resumenes["cwe_por_cve"],
            "por_mes": resumenes["por_mes"],
        }
        rutas.extend(exportar_csv(tablas, carpeta / "powerbi"))
    return rutas
