from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from conftest import cargar
from vigia import procesamiento as proc
from vigia.exportar import exportar
from vigia.fuentes.cisa_kev import parsear_kev
from vigia.fuentes.github_advisory import parsear_advisory
from vigia.fuentes.nvd import parsear_cve

DESDE, HASTA = date(2025, 1, 1), date(2025, 3, 31)


@pytest.fixture
def registros():
    nvd = [parsear_cve(v["cve"]) for v in cargar("nvd_pagina.json")["vulnerabilities"]]
    ghsa = [r for r in (parsear_advisory(a) for a in cargar("ghsa_pagina.json")) if r]
    kev = [parsear_kev(v) for v in cargar("kev.json")["vulnerabilities"] if v["dateAdded"] >= "2025-01-01"]
    return proc.registros_a_df(nvd + ghsa + kev)


@pytest.fixture
def consolidado(registros):
    return proc.consolidar(registros)


def fila(df, cve):
    return df.set_index("cve_id").loc[cve]


def test_una_fila_por_cve(consolidado):
    assert consolidado["cve_id"].is_unique
    assert set(consolidado["cve_id"]) == {"CVE-2025-0001", "CVE-2025-0002", "CVE-2025-0003",
                                          "CVE-2025-0004", "CVE-2024-5555"}


def test_prioridades_de_consolidacion(consolidado):
    a = fila(consolidado, "CVE-2025-0001")
    assert a["cvss_puntaje"] == 9.8                  # CVSS de la NVD
    assert a["proveedor"] == "Example Corp"           # proveedor de KEV
    assert a["titulo"].startswith("Example Corp Web Server OS")  # título de KEV
    assert a["fuentes"] == "NVD, CISA KEV"
    assert a["explotada_activamente"] and a["ransomware"] == "Sí"
    assert a["fecha_limite_cisa"] == pd.Timestamp("2025-01-31")
    assert a["url"].startswith("https://nvd.nist.gov")

    b = fila(consolidado, "CVE-2025-0003")           # NVD sin CVSS + GitHub con CVSS
    assert b["cvss_puntaje"] == 5.6 and b["severidad"] == "MEDIA"
    assert b["en_nvd"] and b["en_github"] and not b["en_kev"]
    assert b["titulo"] == "Prototype pollution in example-lib"

    c = fila(consolidado, "CVE-2024-5555")           # solo KEV
    assert c["severidad"] == "SIN CALIFICAR" and c["severidad_orden"] == 0


def test_tablas_resumen(registros, consolidado):
    por_mes = proc.tabla_por_mes(registros, consolidado, DESDE, HASTA, ["NVD", "CISA KEV", "GitHub Advisory"])
    assert list(por_mes["Mes"]) == ["2025-01", "2025-02", "2025-03"]
    assert por_mes.loc[por_mes["Mes"] == "2025-03", "GitHub Advisory"].iat[0] == 2
    sev = proc.tabla_por_severidad(consolidado)
    assert sev["Cantidad"].sum() == len(consolidado)
    top = proc.tabla_top_cwe(consolidado)
    assert "CWE-78" in set(top["CWE"])
    kpis = proc.indicadores(consolidado, registros)
    assert kpis["CVE únicas"] == 5 and kpis["Usadas en ransomware"] == 1


def test_texto_seguro_quita_caracteres_que_rompen_excel():
    from vigia.exportar import _texto_seguro
    assert _texto_seguro("abc\uffff\x07def") == "abcdef"
    assert _texto_seguro("=1+1") == " =1+1"


def test_csv_una_fila_por_linea(tmp_path):
    from vigia.exportar import exportar_csv
    df = pd.DataFrame({"cve_id": ["CVE-1", "CVE-2"], "descripcion": ["línea 1\nlínea 2\r\nfin", "normal"]})
    ruta = exportar_csv({"t": df}, tmp_path)[0]
    lineas = ruta.read_text(encoding="utf-8-sig").splitlines()
    assert len(lineas) == 3 and "línea 1 línea 2 fin" in lineas[1]


def test_limpiar_tambien_columnas_str():
    """En pandas 3 el texto es dtype 'str'; la protección debe aplicarse igual."""
    from vigia.exportar import limpiar
    df = pd.DataFrame({"t": pd.array(["=HYPERLINK(1)", "ok\uffff"], dtype="string")})
    limpio = limpiar(df)
    assert list(limpio["t"]) == [" =HYPERLINK(1)", "ok"]


def test_consolidar_vacio():
    assert proc.consolidar(proc.registros_a_df([])).empty


def test_exportar_excel_y_csv(tmp_path, registros, consolidado):
    resumenes = {
        "por_mes": proc.tabla_por_mes(registros, consolidado, DESDE, HASTA),
        "por_severidad": proc.tabla_por_severidad(consolidado),
        "top_proveedores": proc.tabla_top_proveedores(consolidado),
        "top_cwe": proc.tabla_top_cwe(consolidado),
        "cwe_por_cve": proc.tabla_cwe_por_cve(consolidado),
    }
    kpis = proc.indicadores(consolidado, registros)
    parametros = {"Generado": "hoy", "Desde": str(DESDE), "Hasta": str(HASTA), "Fuentes consultadas": "todas"}
    rutas = exportar(tmp_path, "ambos", DESDE, HASTA, registros, consolidado, resumenes, kpis, parametros)

    excel = next(r for r in rutas if r.suffix == ".xlsx")
    wb = load_workbook(excel)
    assert wb.sheetnames == ["Resumen", "Consolidado", "Registros", "Por_Mes", "Por_Severidad",
                             "Top_Proveedores", "Top_CWE", "Tendencia", "Parametros"]
    assert len(wb["Resumen"]._charts) == 4
    assert "tblConsolidado" in wb["Consolidado"].tables

    # Protección contra inyección de fórmulas: ninguna celda de texto debe empezar con '='
    for fila_excel in wb["Registros"].iter_rows(min_row=2, values_only=True):
        for valor in fila_excel:
            assert not (isinstance(valor, str) and valor.startswith("="))

    csv = pd.read_csv(tmp_path / "powerbi" / "consolidado.csv", encoding="utf-8-sig")
    assert len(csv) == 5
    assert {"cwe_por_cve.csv", "registros.csv", "por_mes.csv"} <= {p.name for p in (tmp_path / "powerbi").iterdir()}


def test_graficos_no_tapan_el_area(tmp_path, registros, consolidado):
    """En Excel, sin <c:overlay val="0"/> el título y la leyenda se montan encima del gráfico."""
    import re
    import zipfile

    from vigia.exportar import exportar_excel

    resumenes = {
        "por_mes": proc.tabla_por_mes(registros, consolidado, DESDE, HASTA),
        "por_severidad": proc.tabla_por_severidad(consolidado),
        "top_proveedores": proc.tabla_top_proveedores(consolidado),
        "top_cwe": proc.tabla_top_cwe(consolidado),
    }
    parametros = {"Generado": "hoy", "Desde": str(DESDE), "Hasta": str(HASTA), "Fuentes consultadas": "todas"}
    ruta = exportar_excel(tmp_path / "r.xlsx", registros, consolidado, resumenes,
                          proc.indicadores(consolidado, registros), parametros)
    with zipfile.ZipFile(ruta) as z:
        graficos = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        assert len(graficos) == 4
        for nombre in graficos:
            xml = z.read(nombre).decode()
            titulo = re.search(r"<title>.*?</title>", xml, re.S).group(0)
            sin_encima = re.compile(r'<overlay val="0"\s*/>')  # con o sin espacio según la librería XML
            assert sin_encima.search(titulo), nombre
            for leyenda in re.findall(r"<legend>.*?</legend>", xml, re.S):
                assert sin_encima.search(leyenda), nombre
